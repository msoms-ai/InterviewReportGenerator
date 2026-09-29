import zipfile
import re
import os
import docx
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

source_file = r"c:\AntiGravity_Apps\Interview_Feedback\Updated Interview Report Form.docx"
target_file = r"c:\AntiGravity_Apps\Interview_Feedback\platform\public\template.docx"

doc = docx.Document(source_file)

# Fields that are strictly horizontal pairs (Label in cell i, Value in cell i+1)
horizontal_mappings = {
    "Applicant  Name": "+++applicantName+++",
    "Dep/Job Interview For": "+++departmentJob+++",
    "Date": "+++interviewDate+++",
    "Nationality": "+++nationality+++",
    "Date of Birth": "+++dateOfBirth+++",
    "Qualification": "+++qualification+++",
    "Marital Status": "+++maritalStatus+++",
    "Most Recent Job Held": "+++mostRecentJob+++",
    "Year of Relevant Experience": "+++yearsOfExperience+++",
}

# Fields that are full-width cells (Label is at top of cell, Value should go below it)
vertical_in_cell_mappings = {
    "Professional (e.g. academic qualification, technical training, and experience)": "+++professionalComments+++",
    "Personality (e.g. behaviour, attitude, and presentation)": "+++personalityComments+++",
    "Other Comments": "+++otherComments+++",
}

bottom_headers = {
    "Interviewer’s Name": "+++interviewerName+++",
    "Emp #": "+++empNumber+++",
    "Signature": "+++%signatureImage+++",
    "Title": "+++title+++"
}

ratings_map = {
    "Appearance & Confidence": "app",
    "Planning & Execution": "plan",
    "Interpersonal & Communication Skills": "inter",
    "Customer Centricity": "cust",
    "Collaboration": "collab",
    "Agility": "agil",
    "Empowerment": "emp",
    "Relevant Work & Training Experience": "rel",
    "OVERALL RATING": "overall"
}

from docx.shared import Pt

def set_styled_tag(cell, tag_text):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(tag_text)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 255)

def append_styled_tag(cell, tag_text):
    # Remove all empty paragraphs to save vertical space (user likely pressed Enter for handwriting space)
    for p in list(cell.paragraphs):
        if not p.text.strip() and p != cell.paragraphs[0]:
            p._element.getparent().remove(p._element)
            
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(tag_text)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 255)

for table in doc.tables:
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            # Because merged cells are repeated in row.cells, we keep track to avoid duplicating
            # We can use the cell's ID or just check if we already added the tag
            text = cell.text.strip()
            
            # Standard horizontal pairs
            if text in horizontal_mappings:
                if c_idx + 1 < len(row.cells):
                    target = row.cells[c_idx+1]
                    if target != cell and horizontal_mappings[text] not in target.text:
                        set_styled_tag(target, horizontal_mappings[text])
            
            # Vertical in-cell pairs (like Professional, Personality)
            if text in vertical_in_cell_mappings:
                if vertical_in_cell_mappings[text] not in cell.text:
                    append_styled_tag(cell, vertical_in_cell_mappings[text])
            
            # Bottom signature section (vertical pairs)
            if "Interviewer" in text and "Name" in text and c_idx == 0:
                text_key = "Interviewer’s Name"
            else:
                text_key = text

            if text_key in bottom_headers:
                if r_idx + 1 < len(table.rows):
                    target_cell = table.rows[r_idx+1].cells[c_idx]
                    if bottom_headers[text_key] not in target_cell.text:
                        set_styled_tag(target_cell, bottom_headers[text_key])
            
            # Ratings Grid
            if text in ratings_map:
                prefix = ratings_map[text]
                if len(row.cells) >= 7:
                    if f"+++{prefix}_L+++" not in row.cells[2].text:
                        set_styled_tag(row.cells[2], f"+++{prefix}_L+++")
                        set_styled_tag(row.cells[3], f"+++{prefix}_A+++")
                        set_styled_tag(row.cells[4], f"+++{prefix}_G+++")
                        set_styled_tag(row.cells[5], f"+++{prefix}_VG+++")
                        set_styled_tag(row.cells[6], f"+++{prefix}_O+++")

            # Grade line
            if "Grade:" in text and "Work Location:" in text:
                if "+++grade+++" not in cell.text:
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    
                    p.add_run("Grade: ")
                    r1 = p.add_run("+++grade+++")
                    r1.bold = True
                    r1.font.color.rgb = RGBColor(0, 0, 255)
                    
                    p.add_run("   Work Location: ")
                    r2 = p.add_run("+++workLocation+++")
                    r2.bold = True
                    r2.font.color.rgb = RGBColor(0, 0, 255)
                    
                    p.add_run("   Available to join from: ")
                    r3 = p.add_run("+++availableToJoinFrom+++")
                    r3.bold = True
                    r3.font.color.rgb = RGBColor(0, 0, 255)
                    
            # Eligibility and Job Recommended For
            if "Eligibility for Employment" in text and "+++eligibilityForEmp+++" not in cell.text:
                append_styled_tag(cell, "+++eligibilityForEmp+++")
                
            if "Job Recommended for" in text and "+++jobRecommendedFor+++" not in cell.text:
                # Add it at the bottom of THIS cell
                append_styled_tag(cell, "+++jobRecommendedFor+++")
                
            if "Any Relatives working in Etisalat" in text and "+++relativesInEtisalat+++" not in cell.text:
                append_styled_tag(cell, "+++relativesInEtisalat+++")
                
            if "If yes, then details of the relative" in text and "+++relativeDetails+++" not in cell.text:
                append_styled_tag(cell, "+++relativeDetails+++")

doc.save(target_file)
print(f"Clean template saved to {target_file}")
